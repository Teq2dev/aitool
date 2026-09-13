# excel_generator.py — OpenPyXL Excel Workbook (.xlsx) Generator with Multi-Sheet Support
"""
Generates genuine Microsoft Excel (.xlsx) workbooks with:
 - Royal Blue (#2563EB) styled header row
 - Zebra striping for improved legibility
 - Auto-adjusted column widths
 - Thin grid borders
 - Automatic numeric and currency type conversion
 - Multi-table & Multi-sheet support (e.g. Purchase, Sale, Stock sheets)
"""

import io
import re
from typing import List, Dict, Any, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def format_worksheet(ws, headers: List[str], rows: List[List[str]]):
    ws.views.sheetView[0].showGridLines = True

    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    data_font = Font(name="Segoe UI", size=10.5, color="1E293B")
    data_fill_even = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    data_fill_odd = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )
    header_border = Border(
        left=Side(style='thin', color='1D4ED8'),
        right=Side(style='thin', color='1D4ED8'),
        top=Side(style='thin', color='1D4ED8'),
        bottom=Side(style='medium', color='1E40AF')
    )

    # 1. Write Header Row
    if headers and any(h.strip() for h in headers):
        ws.append(headers)
        ws.row_dimensions[1].height = 28
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
            cell.border = header_border

    # 2. Write Data Rows
    start_row = 2 if (headers and any(h.strip() for h in headers)) else 1
    for r_idx, row_data in enumerate(rows, start=start_row):
        typed_row = []
        for val in row_data:
            s_val = str(val or "").strip()
            # Clean currency / comma formatting for numeric check
            clean_s = s_val.replace("$", "").replace(",", "")
            if clean_s and (clean_s.replace(".", "", 1).replace("-", "", 1).isdigit()):
                try:
                    typed_row.append(float(clean_s) if "." in clean_s else int(clean_s))
                except ValueError:
                    typed_row.append(s_val)
            else:
                typed_row.append(s_val)

        ws.append(typed_row)
        ws.row_dimensions[r_idx].height = 22
        is_odd = (r_idx % 2 == 1)
        fill = data_fill_odd if is_odd else data_fill_even

        for col_idx in range(1, len(typed_row) + 1):
            cell = ws.cell(row=r_idx, column=col_idx)
            cell.font = data_font
            cell.fill = fill
            cell.border = thin_border

            if isinstance(cell.value, (int, float)):
                cell.alignment = Alignment(horizontal="right", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=False)

    # 3. Auto-calculate Column Widths
    for col in ws.columns:
        max_len = 10
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(45, max(12, max_len + 3))


def generate_xlsx(
    headers: Optional[List[str]] = None,
    rows: Optional[List[List[str]]] = None,
    sheet_name: str = "Table Data",
    tables: Optional[List[Dict[str, Any]]] = None
) -> bytes:
    """
    Generate a genuine formatted Microsoft Excel (.xlsx) workbook.
    If `tables` is provided (multi-table), creates separate sheets for each table.
    Otherwise, creates a single sheet with `headers` and `rows`.
    """
    wb = openpyxl.Workbook()

    if tables and len(tables) > 1:
        # Multi-sheet workbook
        wb.remove(wb.active) # remove default sheet
        for idx, tbl in enumerate(tables):
            name = tbl.get("name") or f"Table {idx+1}"
            safe_name = re.sub(r'[:\\/\?\*\[\]]', '', name)[:31]
            ws = wb.create_sheet(title=safe_name)
            format_worksheet(ws, tbl.get("headers", []), tbl.get("rows", []))
    else:
        # Single sheet workbook
        ws = wb.active
        tbl_headers = headers or (tables[0].get("headers", []) if tables else [])
        tbl_rows = rows or (tables[0].get("rows", []) if tables else [])
        title = sheet_name
        if tables and tables[0].get("name"):
            title = tables[0]["name"]
        safe_title = re.sub(r'[:\\/\?\*\[\]]', '', title)[:31]
        ws.title = safe_title
        format_worksheet(ws, tbl_headers, tbl_rows)

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
