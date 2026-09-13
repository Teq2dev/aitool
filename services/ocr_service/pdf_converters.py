"""
pdf_converters.py — Production-grade document conversion routines
Supported Conversions:
- PDF -> Word (.docx) via pdf2docx + PyMuPDF fallback
- PDF -> PowerPoint (.pptx) via PyMuPDF + python-pptx
- PDF -> Excel (.xlsx) via pdfplumber + openpyxl
- Word (.docx) -> PDF via python-docx + reportlab
- PowerPoint (.pptx) -> PDF via python-pptx + reportlab
- Excel (.xlsx) -> PDF via openpyxl + reportlab
"""

import io
import os
import tempfile
import re
from typing import Optional

def pdf_to_docx(pdf_bytes: bytes) -> bytes:
    """Converts PDF document bytes into an editable Microsoft Word (.docx) document."""
    temp_pdf = None
    temp_docx = None
    try:
        from pdf2docx import Converter
        with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as f_pdf:
            f_pdf.write(pdf_bytes)
            temp_pdf = f_pdf.name
        temp_docx = temp_pdf.replace(".pdf", ".docx")
        
        cv = Converter(temp_pdf)
        cv.convert(temp_docx, start=0, end=None)
        cv.close()
        
        with open(temp_docx, "rb") as f_out:
            return f_out.read()
    except Exception as e:
        import fitz
        from docx import Document
        from docx.shared import Inches
        doc = Document()
        pdf_doc = fitz.open(stream=pdf_bytes, filetype="pdf")
        for page_idx, page in enumerate(pdf_doc):
            if page_idx > 0:
                doc.add_page_break()
            text = page.get_text("text")
            if text.strip():
                for para in text.split("\n\n"):
                    clean = para.strip()
                    if clean:
                        doc.add_paragraph(clean)
            else:
                pix = page.get_pixmap(dpi=150)
                img_bytes = pix.tobytes("png")
                doc.add_picture(io.BytesIO(img_bytes), width=Inches(6.0))
                
        out_stream = io.BytesIO()
        doc.save(out_stream)
        return out_stream.getvalue()
    finally:
        if temp_pdf and os.path.exists(temp_pdf):
            try: os.remove(temp_pdf)
            except Exception: pass
        if temp_docx and os.path.exists(temp_docx):
            try: os.remove(temp_docx)
            except Exception: pass

def pdf_to_pptx(pdf_bytes: bytes) -> bytes:
    """Converts PDF document pages into high-fidelity PowerPoint slides (.pptx)."""
    import fitz
    from pptx import Presentation
    from pptx.util import Inches

    pdf_doc = fitz.open(stream=pdf_bytes, filetype="pdf")
    prs = Presentation()
    blank_layout = prs.slide_layouts[6]
    
    for page in pdf_doc:
        rect = page.rect
        width_in = max(rect.width / 72.0, 4.0)
        height_in = max(rect.height / 72.0, 3.0)
        prs.slide_width = Inches(width_in)
        prs.slide_height = Inches(height_in)
        
        slide = prs.slides.add_slide(blank_layout)
        pix = page.get_pixmap(dpi=200)
        img_bytes = pix.tobytes("png")
        
        slide.shapes.add_picture(
            io.BytesIO(img_bytes),
            Inches(0),
            Inches(0),
            width=Inches(width_in),
            height=Inches(height_in)
        )

    out_stream = io.BytesIO()
    prs.save(out_stream)
    return out_stream.getvalue()

def pdf_to_xlsx(pdf_bytes: bytes) -> bytes:
    """Extracts tables and tabular rows from PDF pages into formatted Excel (.xlsx) workbook."""
    import pdfplumber
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid")
    regular_font = Font(name="Arial", size=10)
    border_side = Side(border_style="thin", color="E2E8F0")
    thin_border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)
    center_align = Alignment(vertical="center")

    has_data = False
    with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
        for idx, page in enumerate(pdf.pages):
            sheet_title = f"Page {idx + 1}"
            ws = wb.create_sheet(title=sheet_title)
            
            tables = page.extract_tables()
            current_row = 1

            if tables:
                for table in tables:
                    if not table:
                        continue
                    has_data = True
                    for r_idx, row in enumerate(table):
                        if not any(row):
                            continue
                        for c_idx, val in enumerate(row):
                            cell = ws.cell(row=current_row, column=c_idx + 1, value=str(val or "").strip())
                            cell.border = thin_border
                            cell.alignment = center_align
                            if r_idx == 0:
                                cell.font = header_font
                                cell.fill = header_fill
                            else:
                                cell.font = regular_font
                        current_row += 1
                    current_row += 2
            else:
                text = page.extract_text()
                if text:
                    has_data = True
                    for line in text.split("\n"):
                        clean = line.strip()
                        if not clean:
                            continue
                        parts = re.split(r"\s{2,}|\t+", clean)
                        for c_idx, part in enumerate(parts):
                            cell = ws.cell(row=current_row, column=c_idx + 1, value=part.strip())
                            cell.font = regular_font
                            cell.alignment = center_align
                        current_row += 1

            for col in ws.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    if cell.value:
                        max_len = max(max_len, len(str(cell.value)))
                ws.column_dimensions[col_letter].width = max(max_len + 3, 10)

    if not has_data or len(wb.sheetnames) == 0:
        ws = wb.create_sheet(title="Data")
        ws.cell(row=1, column=1, value="No tabular data extracted from PDF.")

    out_stream = io.BytesIO()
    wb.save(out_stream)
    return out_stream.getvalue()

def docx_to_pdf(docx_bytes: bytes) -> bytes:
    """Converts Word document (.docx) bytes into a formatted PDF document."""
    from docx import Document
    from reportlab.lib.pagesizes import letter
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

    doc = Document(io.BytesIO(docx_bytes))
    out_stream = io.BytesIO()
    pdf_doc = SimpleDocTemplate(
        out_stream,
        pagesize=letter,
        leftMargin=54,
        rightMargin=54,
        topMargin=54,
        bottomMargin=54
    )

    styles = getSampleStyleSheet()
    normal_style = styles["Normal"]
    normal_style.fontSize = 10
    normal_style.leading = 14
    
    h1_style = styles["Heading1"]
    h1_style.fontSize = 18
    h1_style.leading = 22
    h1_style.textColor = colors.HexColor("#0F172A")

    h2_style = styles["Heading2"]
    h2_style.fontSize = 14
    h2_style.leading = 18
    h2_style.textColor = colors.HexColor("#1E293B")

    flowables = []

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            flowables.append(Spacer(1, 8))
            continue
        
        style_name = para.style.name.lower() if para.style else ""
        if "heading 1" in style_name or "title" in style_name:
            flowables.append(Paragraph(text, h1_style))
            flowables.append(Spacer(1, 10))
        elif "heading 2" in style_name:
            flowables.append(Paragraph(text, h2_style))
            flowables.append(Spacer(1, 8))
        else:
            flowables.append(Paragraph(text, normal_style))
            flowables.append(Spacer(1, 6))

    for table in doc.tables:
        table_data = []
        for row in table.rows:
            row_data = [Paragraph(cell.text.strip(), normal_style) for cell in row.cells]
            table_data.append(row_data)
        if table_data:
            t = Table(table_data, colWidths=None)
            t.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F1F5F9")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#0F172A")),
                ("ALIGN", (0, 0), (-1, -1), "LEFT"),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
            ]))
            flowables.append(t)
            flowables.append(Spacer(1, 12))

    if not flowables:
        flowables.append(Paragraph("Empty Document", normal_style))

    pdf_doc.build(flowables)
    return out_stream.getvalue()

def pptx_to_pdf(pptx_bytes: bytes) -> bytes:
    """Converts PowerPoint presentation (.pptx) bytes into a multi-page PDF document."""
    from pptx import Presentation
    from reportlab.pdfgen import canvas
    from reportlab.lib import colors

    prs = Presentation(io.BytesIO(pptx_bytes))
    width_pt = prs.slide_width.pt if prs.slide_width else 720
    height_pt = prs.slide_height.pt if prs.slide_height else 540

    out_stream = io.BytesIO()
    c = canvas.Canvas(out_stream, pagesize=(width_pt, height_pt))

    for slide in prs.slides:
        c.setFillColor(colors.white)
        c.rect(0, 0, width_pt, height_pt, fill=1, stroke=0)

        for shape in slide.shapes:
            if shape.has_text_frame:
                tf = shape.text_frame
                left = shape.left.pt
                top = shape.top.pt
                pdf_y = height_pt - top
                
                c.setFillColor(colors.HexColor("#0F172A"))
                y_offset = pdf_y - 14
                for p in tf.paragraphs:
                    line_text = p.text.strip()
                    if line_text:
                        font_size = p.font.size.pt if p.font and p.font.size else 12
                        c.setFont("Helvetica", font_size)
                        c.drawString(left, max(y_offset, 10), line_text)
                        y_offset -= (font_size + 4)
            
            if shape.shape_type == 13:
                try:
                    img_bytes = shape.image.blob
                    img_io = io.BytesIO(img_bytes)
                    img_left = shape.left.pt
                    img_top = shape.top.pt
                    img_w = shape.width.pt
                    img_h = shape.height.pt
                    pdf_y = height_pt - img_top - img_h
                    from reportlab.lib.utils import ImageReader
                    c.drawImage(ImageReader(img_io), img_left, pdf_y, width=img_w, height=img_h)
                except Exception:
                    pass

        c.showPage()

    c.save()
    return out_stream.getvalue()

def xlsx_to_pdf(xlsx_bytes: bytes) -> bytes:
    """Converts Excel spreadsheet (.xlsx) bytes into formatted paginated PDF tables."""
    import openpyxl
    from reportlab.lib.pagesizes import letter, landscape
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), data_only=True)
    out_stream = io.BytesIO()
    pdf_doc = SimpleDocTemplate(
        out_stream,
        pagesize=landscape(letter),
        leftMargin=36,
        rightMargin=36,
        topMargin=36,
        bottomMargin=36
    )

    styles = getSampleStyleSheet()
    title_style = styles["Heading2"]
    title_style.textColor = colors.HexColor("#0F172A")
    cell_style = styles["Normal"]
    cell_style.fontSize = 8
    cell_style.leading = 10

    flowables = []

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        flowables.append(Paragraph(f"<b>Sheet: {sheet_name}</b>", title_style))
        flowables.append(Spacer(1, 10))

        rows_data = []
        max_cols = 0
        for row in ws.iter_rows(values_only=True):
            if any(row):
                str_row = [Paragraph(str(val) if val is not None else "", cell_style) for val in row]
                rows_data.append(str_row)
                max_cols = max(max_cols, len(str_row))

        if rows_data:
            for r in rows_data:
                while len(r) < max_cols:
                    r.append(Paragraph("", cell_style))
            
            avail_w = 720
            col_w = max(avail_w / max_cols, 35) if max_cols > 0 else 50
            col_widths = [col_w] * max_cols

            t = Table(rows_data, colWidths=col_widths)
            t.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E3A8A")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
            ]))
            flowables.append(t)
        else:
            flowables.append(Paragraph("Empty Worksheet", cell_style))

        flowables.append(Spacer(1, 20))

    if not flowables:
        flowables.append(Paragraph("Empty Workbook", cell_style))

    pdf_doc.build(flowables)
    return out_stream.getvalue()
